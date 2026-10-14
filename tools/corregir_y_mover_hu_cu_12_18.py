import os
import glob
import docx

folder_unrev = r"C:\Users\Andres\Desktop\Personal\Universidad de Cundinamarca\V Semestre\Ingeniería de Software\Historias de Usuario y Casos de Uso\documentos sin revisar generado en claude SYSQUEST"
folder_dest = r"C:\Users\Andres\Desktop\Personal\Universidad de Cundinamarca\V Semestre\Ingeniería de Software\Historias de Usuario y Casos de Uso"

modificaciones = {
    "HU-12": [
        ("Se acordó interpretar RF-12 (\"competencia individual y desempeño\") como un ranking/leaderboard.",
         "Criterio cerrado: Ranking global por xp_total acumulado. No expone datos sensibles ni contraseñas; únicamente nombre_usuario, nivel y xp_total. Permite destacar la posición relativa del estudiante autenticado."),
        ("Pendiente", "Cerrado")
    ],
    "CU-12": [
        ("Pendiente de definir", "< 1.5 segundos"),
        ("Pendiente", "Implementado")
    ],
    "HU-13": [
        ("Pendiente de definir el valor de fuente_generacion para contenido creado por docentes.",
         "Valor oficial definido: enum fuente_generacion = 'DOCENTE' (junto a 'AI' y 'FALLBACK'). Las quests creadas por docentes quedan versionadas e identificadas con este valor."),
        ("Pendiente de definir", "fuente_generacion = 'DOCENTE'"),
        ("Pendiente", "Cerrado")
    ],
    "CU-13": [
        ("Pendiente de definir", "fuente_generacion = 'DOCENTE'"),
        ("Pendiente", "Cerrado")
    ],
    "HU-14": [
        ("Pendiente de confirmar si se restringe la edición de una quest con partidas en curso asociadas.",
         "Regla de negocio oficial cerrada: No se edita en caliente. La edición genera una nueva versión inmutable (version + 1) e id_quest_original vinculante. Las partidas activas continúan en la versión original sin corromper datos."),
        ("Pendiente de confirmar", "Versionado inmutable (version + 1)"),
        ("Pendiente", "Cerrado")
    ],
    "CU-14": [
        ("Pendiente de confirmar si se restringe la edición", "Versionado inmutable: se genera nueva versión de quest sin modificar las partidas en curso"),
        ("Pendiente de definir", "Transaccional / Inmutable"),
        ("Pendiente", "Cerrado")
    ],
    "HU-15": [
        ("Pendiente de confirmar si se requiere la posibilidad de reactivar una quest previamente desactivada.",
         "Regla de negocio oficial cerrada: Borrado lógico obligatorio mediante flag 'activa = false'. Las partidas previas permanecen auditables. Reactivar la quest simplemente restablece 'activa = true' sin eliminar registros físicos."),
        ("Pendiente de confirmar", "Borrado lógico mediante activa = false"),
        ("Pendiente", "Cerrado")
    ],
    "CU-15": [
        ("Pendiente de confirmar", "Borrado lógico (activa = false)"),
        ("Pendiente de definir", "< 1.0 segundo"),
        ("Pendiente", "Cerrado")
    ],
    "HU-16": [
        ("Pendiente", "Cerrado"),
        ("Sin registros adicionales por el momento.",
         "Métricas calculadas en tiempo real mediante consultas SQL optimizadas sobre PARTIDA y PROGRESO_USUARIO: tasa de victoria/derrota, tiempo promedio y encuentros con mayor tasa de error.")
    ],
    "CU-16": [
        ("Pendiente de definir", "< 2.0 segundos"),
        ("Pendiente", "Cerrado")
    ],
    "HU-17": [
        ("Pendiente de definir", "Roles controlados vía Supabase Auth y políticas RLS: 'estudiante', 'docente', 'administrador'"),
        ("Pendiente", "Cerrado"),
        ("Sin registros adicionales por el momento.",
         "Control de acceso gobernado por Supabase Auth (GoTrue). Las cuentas desactivadas tienen su token revocado y se bloquea el acceso en el Gateway.")
    ],
    "CU-17": [
        ("Pendiente de definir", "Políticas RLS en PostgreSQL"),
        ("Pendiente", "Cerrado")
    ],
    "HU-18": [
        ("Pendiente de definir con precisión qué parámetros son configurables (RF-11 no detalla el alcance).",
         "Parámetros configurables oficiales: 1) timeout_ia (10 segundos), 2) xp_victoria (+50), 3) xp_derrota (+10), 4) max_wildcards_inventario (3 unidades por tipo), 5) flag_usar_backend_remoto."),
        ("Pendiente de definir qué parámetros específicos son configurables.",
         "Parámetros definidos: timeout_ia (10s), xp_victoria (50), xp_derrota (10), límite de wildcards (3)."),
        ("Pendiente de definir", "timeout_ia, xp_victoria, xp_derrota, max_wildcards"),
        ("Pendiente", "Cerrado")
    ],
    "CU-18": [
        ("Pendiente de definir qué parámetros específicos", "Parámetros del sistema: timeout IA, XP de partida, límites de comodines"),
        ("Pendiente de definir", "Inmediata en tabla configuracion_sistema"),
        ("Pendiente", "Cerrado")
    ]
}

def procesar_archivo(codigo, tipo):
    prefix = f"{tipo}-{codigo}"
    patron = os.path.join(folder_unrev, f"{prefix}*.docx")
    archivos = [f for f in glob.glob(patron) if "Repaired" not in f]
    if not archivos:
        archivos = glob.glob(patron)
    if not archivos:
        print(f"No se encontró archivo para {prefix}")
        return
    
    archivo_origen = archivos[0]
    nombre_base = os.path.basename(archivo_origen).replace(" (Repaired)", "")
    archivo_destino = os.path.join(folder_dest, nombre_base)
    
    doc = docx.Document(archivo_origen)
    modos = modificaciones.get(prefix, [])
    
    cambios_realizados = 0
    # Reemplazar en párrafos
    for p in doc.paragraphs:
        for viejo, nuevo in modos:
            if viejo in p.text:
                p.text = p.text.replace(viejo, nuevo)
                cambios_realizados += 1
                
    # Reemplazar en tablas
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for viejo, nuevo in modos:
                    if viejo in cell.text:
                        cell.text = cell.text.replace(viejo, nuevo)
                        cambios_realizados += 1

    doc.save(archivo_destino)
    print(f"OK: {nombre_base} corregido y copiado ({cambios_realizados} correcciones aplicadas)")

for num in ["12", "13", "14", "15", "16", "17", "18"]:
    procesar_archivo(num, "HU")
    procesar_archivo(num, "CU")

print("Todas las HU y CU de 12 a 18 han sido procesadas con éxito.")

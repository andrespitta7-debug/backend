import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

source_path = r"C:\Users\Andres\Desktop\Personal\Universidad de Cundinamarca\V Semestre\Ingeniería de Software\SysQuest - Ing de Software  V1.13.docx"
target_path = r"C:\Users\Andres\Desktop\Personal\Universidad de Cundinamarca\V Semestre\Ingeniería de Software\SysQuest - Ing de Software  V1.14 - Comite.docx"
image_path = r"C:\Users\Andres\.gemini\antigravity\brain\9c097cc1-aee5-4498-a8d9-c72b349a7bed\.user_uploaded\media_1791490959050.png"

doc = docx.Document(source_path)

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def format_cell_text(cell, text, bold=False, italic=False, font_size=9, color_rgb=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="6" w:space="0" w:color="003366"/>
        <w:bottom w:val="single" w:sz="6" w:space="0" w:color="003366"/>
        <w:left w:val="none"/>
        <w:right w:val="none"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def add_header(title, level=1):
    h = doc.add_paragraph()
    h.style = doc.styles['Heading 1'] if level == 1 else doc.styles['Heading 2']
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(title)
    run.font.name = 'Calibri'
    run.font.size = Pt(14 if level == 1 else 12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102) # Navy Blue
    return h

def add_body_paragraph(text, bold_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 30, 30)
    r_body = p.add_run(text)
    r_body.font.name = 'Calibri'
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = RGBColor(40, 40, 40)
    return p

# ----------------- SECCION 10: ARQUITECTURA DETALLADA -----------------
add_header("10. Arquitectura de Software y Diagrama de Componentes", level=1)
add_body_paragraph(
    "Para garantizar la separación estricta de responsabilidades, la alta disponibilidad y la capacidad de evolución técnica, "
    "SysQuest implementa una arquitectura desacoplada basada en Clean Architecture en la capa cliente (Flutter) y una infraestructura "
    "serverless en el borde (Edge Computing) provista por Supabase. A continuación se presenta el diagrama de componentes detallado "
    "de la solución, ilustrando los canales de comunicación, subsistemas y dependencias de persistencia e inteligencia artificial."
)

p_fig = doc.add_paragraph()
p_fig.paragraph_format.space_before = Pt(6)
p_fig.paragraph_format.space_after = Pt(2)
r_fig_title = p_fig.add_run("Figura 3\nDiagrama de Componentes Detallado de SysQuest")
r_fig_title.font.name = 'Calibri'
r_fig_title.font.size = Pt(10)
r_fig_title.font.bold = True

if os.path.exists(image_path):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(4)
    p_img.paragraph_format.space_after = Pt(4)
    p_img.add_run().add_picture(image_path, width=Inches(6.2))

p_nota = doc.add_paragraph()
p_nota.paragraph_format.space_before = Pt(2)
p_nota.paragraph_format.space_after = Pt(8)
r_nota = p_nota.add_run("Nota. Diagrama oficial de arquitectura y componentes de SysQuest. Elaboración propia.")
r_nota.font.name = 'Calibri'
r_nota.font.size = Pt(8.5)
r_nota.font.italic = True

add_body_paragraph(
    "1. Capa Cliente (Frontend - Flutter): ", bold_prefix="• "
)
add_body_paragraph(
    "Estructurada bajo principios de Arquitectura Limpia, dividiendo la aplicación en módulos funcionales: Autenticación, Gestión de Quests, "
    "Encuentros y Combate, Progreso y Estadísticas, Personaje, Administración y Configuración. La interfaz interactúa con la lógica de negocio "
    "a través de controladores reactivos y proveedores de estado (Provider), aislando totalmente las pantallas de la tecnología de datos."
)
add_body_paragraph(
    "2. Capa Servidora (Backend - Supabase Edge Functions): ", bold_prefix="• "
)
add_body_paragraph(
    "Desplegada en contenedores efímeros bajo el motor de ejecución Deno en el borde (Edge). Expone endpoints REST/HTTPS protegidos para "
    "gestionar usuarios, partidas, generación procedural de quests y cálculo de recompensas. El cliente Flutter nunca interactúa de manera directa "
    "con la base de datos relacional, canalizando toda operación a través de esta capa intermedia de servicios."
)
add_body_paragraph(
    "3. Persistencia Híbrida y Servicios de Soporte: ", bold_prefix="• "
)
add_body_paragraph(
    "El sistema cuenta con un modelo de persistencia dual: persistencia local en SQLite (garantizando operabilidad offline mediante repositorios stub) "
    "y persistencia remota centralizada en PostgreSQL administrada por Supabase. La autenticación segura se delega al microservicio Supabase Auth (GoTrue) "
    "mediante JSON Web Tokens (JWT), y la generación narrativa e interactiva se delega a proveedores externos de Inteligencia Artificial (Gemini y Groq)."
)

# ----------------- SECCION 11: METODOLOGIA 12 FACTORES -----------------
add_header("11. Adopción de la Metodología The Twelve-Factor App", level=1)
add_body_paragraph(
    "Con el objetivo de seguir los más altos estándares de la ingeniería de software para el desarrollo de aplicaciones nativas en la nube "
    "(cloud-native) y software como servicio (SaaS), el backend de SysQuest se estructuró bajo los 12 factores de diseño propuestos por Adam Wiggins:"
)

factores = [
    ("Factor 1 - Base de Código Única (Codebase): ", "Un único repositorio rastreado bajo control de versiones (Git/GitHub) del cual se derivan múltiples despliegues controlados (local en Docker y nube en Supabase Edge Runtime)."),
    ("Factor 2 - Dependencias Explícitas (Dependencies): ", "Todas las dependencias de Deno y Dart se declaran de forma explícita e inmutable mediante URLs versionadas e import maps, evitando dependencias implícitas en el sistema anfitrión."),
    ("Factor 3 - Configuración en Variables de Entorno (Config): ", "Separación absoluta entre código y configuración. Parámetros sensibles como SUPABASE_URL, SUPABASE_ANON_KEY y llaves de IA se inyectan en tiempo de ejecución desde el entorno, jamás embebidas en el repositorio."),
    ("Factor 4 - Servicios de Respaldo Desacoplados (Backing Services): ", "La base de datos PostgreSQL, el servicio de autenticación y los proveedores de IA son tratados como recursos vinculados intercambiables mediante cadenas de conexión y credenciales seguras."),
    ("Factor 5 - Separación de Construcción, Distribución y Ejecución (Build, Release, Run): ", "Etapas estrictamente separadas. Cada compilación genera artefactos versionados inmutables y etiquetados, vinculados a migraciones atómicas en SQL."),
    ("Factor 6 - Procesos Sin Estado (Processes): ", "Las Edge Functions operan como procesos totalmente stateless y sin memoria compartida. Toda persistencia requerida se almacena de forma inmediata en PostgreSQL o SQLite."),
    ("Factor 7 - Asignación de Puertos (Port Binding): ", "La aplicación es completamente autónoma; en desarrollo local expone sus servicios en puertos HTTP independientes (54321) mediante el API Gateway Kong embebido."),
    ("Factor 8 - Concurrencia Horizontal (Concurrency): ", "Escalabilidad elástica mediante el modelo de ejecución V8 Isolates de Supabase Edge Runtime, permitiendo atender cientos de solicitudes en paralelo sin sobrecargar recursos del servidor."),
    ("Factor 9 - Desechabilidad (Disposability): ", "Procesos con tiempos de arranque instantáneos (cold starts inferiores a 200 ms) y finalización limpia ante fallos, facilitando la tolerancia a fallos y el reinicio dinámico."),
    ("Factor 10 - Paridad en Desarrollo y Producción (Dev/Prod Parity): ", "Máxima similitud entre entornos mediante contenedores Docker locales que replican con exactitud el comportamiento, versiones y configuraciones de PostgreSQL y Kong en la nube."),
    ("Factor 11 - Registro y Monitoreo como Flujos de Eventos (Logs): ", "Los eventos del sistema nunca se escriben en archivos de texto fijos en el servidor. Se tratan como flujos continuos de eventos estructurados capturados por Logflare y el panel de telemetría de Supabase."),
    ("Factor 12 - Procesos de Administración Idempotentes (Admin Processes): ", "Las tareas de inicialización, siembra de datos de prueba (seeds) y migraciones de esquema se ejecutan como scripts SQL automatizados e independientes del ciclo de vida del cliente.")
]

for tit, desc in factores:
    add_body_paragraph(desc, bold_prefix=tit)

# ----------------- SECCION 12: ESTRATEGIA DE LOGS -----------------
add_header("12. Estrategia de Registro, Auditoría y Gestión de Logs", level=1)
add_body_paragraph(
    "El registro de eventos es un pilar fundamental para la trazabilidad operativa, la detección de incidentes de seguridad y el análisis "
    "del desempeño de los modelos de inteligencia artificial integrados en SysQuest. La estrategia implementada contempla:"
)

add_body_paragraph(
    "1. Categorización de Logs: ", bold_prefix="• "
)
add_body_paragraph(
    "Se manejan tres niveles estrictos: INFO para eventos transaccionales exitosos (inicio de sesión, finalización de combate, otorgamiento de XP); "
    "WARN para anomalías no bloqueantes (tiempos de respuesta de IA superiores a 8 segundos, activación de mecanismos de fallback por indisponibilidad del proveedor); "
    "y ERROR para excepciones no controladas, fallos en inserciones transaccionales o violaciones de esquemas."
)
add_body_paragraph(
    "2. Logs Estructurados en JSON: ", bold_prefix="• "
)
add_body_paragraph(
    "Cada registro emitido por las Edge Functions contiene metadatos contextuales: timestamp en UTC (ISO-8601), request_id, endpoint invocado, "
    "id_usuario (si la sesión está autenticada), código de estado HTTP y latencia en milisegundos. Esta uniformidad permite indexar y filtrar eventos con rapidez."
)
add_body_paragraph(
    "3. Centralización e Ingesta de Eventos: ", bold_prefix="• "
)
add_body_paragraph(
    "A través de la integración de Supabase Analytics y Logflare, los flujos emitidos hacia stdout/stderr son recolectados en tiempo real. "
    "Esto evita la pérdida de trazas por reinicio de contenedores efímeros y permite configurar alertas automáticas ante ráfagas de errores 500 o accesos no autorizados (401/403)."
)
add_body_paragraph(
    "4. Auditoría y Trazabilidad de Contenido Pedagógico por IA: ", bold_prefix="• "
)
add_body_paragraph(
    "Con el propósito de auditar la pertinencia académica y monitorear el consumo de tokens en Groq y Gemini, se almacena una traza de auditoría "
    "que vincula la temática solicitada, el tiempo de inferencia del LLM, el resultado del validador de calidad (V9) y si la misión generada fue aprobada o rechazada."
)

# ----------------- SECCION 13: ATRIBUTOS DE CALIDAD E ISO -----------------
add_header("13. Atributos de Calidad y Alineación con Estándares ISO", level=1)
add_body_paragraph(
    "El diseño arquitectónico y las prácticas de ingeniería adoptadas en SysQuest se fundamentan en las normas internacionales "
    "ISO/IEC 25010 (Modelo de Calidad del Producto de Software) e ISO/IEC 27001 (Sistemas de Gestión de Seguridad de la Información). "
    "A continuación se presenta la matriz de trazabilidad entre los atributos de calidad, su materialización técnica en el proyecto y los mecanismos de verificación establecidos:"
)

# Tabla Atributos de Calidad
table_data = [
    ["Norma", "Atributo de Calidad", "Estrategia e Implementación en SysQuest", "Mecanismo de Verificación"],
    ["ISO/IEC 25010", "Adecuación Funcional", "Cobertura completa del flujo gamificado: registro, generación procedural de quests, encuentros por turnos, sistema de wildcards y cálculo de nivel.", "Batería de 117 pruebas unitarias y de integración automatizadas en Flutter y pruebas en Postman."],
    ["ISO/IEC 25010", "Eficiencia de Desempeño", "Ejecución distribuida en Edge Functions con baja latencia. Algoritmo de caché en memoria para quests recurrentes y optimización de consultas SQL.", "Medición de tiempos de respuesta en Edge Runtime (< 2s) y perfiles de red en Supabase Analytics."],
    ["ISO/IEC 25010", "Usabilidad y Experiencia", "Interfaz retro pixel art con efectos inmersivos de Game Feel (hitstop de 80ms, sacudida de pantalla, animaciones de daño) y navegación fluida.", "Pruebas de usabilidad bajo el RNF-08 (incorporación de nuevos estudiantes en menos de 5 minutos)."],
    ["ISO/IEC 25010", "Fiabilidad y Resiliencia", "Mecanismo de fallback dual: banco de preguntas semilla en PostgreSQL y generador stub local offline en SQLite ante fallos de conectividad.", "Simulación de cortes de red y desconexión de servicios de IA durante la ejecución de partidas."],
    ["ISO/IEC 25010", "Mantenibilidad", "Implementación rigurosa de Clean Architecture (Domain, Infrastructure, Presentation) y Edge Functions modulares e independientes.", "Análisis estático de código mediante el linter oficial de Dart y revisiones periódicas de arquitectura."],
    ["ISO/IEC 25010", "Portabilidad", "Base de código única en Flutter/Dart capaz de compilar nativamente a Android, iOS, Web y sistemas de escritorio.", "Compilación y validación en emuladores móviles y entornos web locales."],
    ["ISO/IEC 27001", "Seguridad: Autenticación", "Delegación del control de identidades a Supabase Auth (GoTrue). Uso de JSON Web Tokens (JWT) con expiración controlada en cada solicitud.", "Verificación de respuestas 401 Unauthorized ante invocaciones sin token en endpoints protegidos."],
    ["ISO/IEC 27001", "Seguridad: Control de Acceso", "Aislamiento granular mediante directivas Row Level Security (RLS) en PostgreSQL; cada usuario solo puede ver y editar sus propios registros.", "Auditoría directa en la base de datos validando que consultas cruzadas entre usuarios son bloqueadas."],
    ["ISO/IEC 27001", "Seguridad: Integridad de Datos", "Operaciones transaccionales atómicas mediante funciones almacenadas RPC (crear_run_completa, finalizar_partida) con restricciones FK activas.", "Pruebas de integridad transaccional simulando desconexiones a mitad de la persistencia de partidas."]
]

table = doc.add_table(rows=len(table_data), cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(table)

for row_idx, row_content in enumerate(table_data):
    row = table.rows[row_idx]
    is_header = (row_idx == 0)
    for col_idx, cell_text in enumerate(row_content):
        cell = row.cells[col_idx]
        if is_header:
            set_cell_shading(cell, "003366") # Navy
            format_cell_text(cell, cell_text, bold=True, font_size=9, color_rgb=RGBColor(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)
        else:
            if row_idx % 2 == 1:
                set_cell_shading(cell, "F4F6F9") # Soft grey/blue zebra
            format_cell_text(cell, cell_text, bold=(col_idx < 2), font_size=8.5, color_rgb=RGBColor(30, 30, 30))

p_post_tbl = doc.add_paragraph()
p_post_tbl.paragraph_format.space_before = Pt(4)
p_post_tbl.paragraph_format.space_after = Pt(8)
r_nota_tbl = p_post_tbl.add_run("Nota. Matriz de correspondencia entre estándares de calidad de software y características implementadas en SysQuest.")
r_nota_tbl.font.name = 'Calibri'
r_nota_tbl.font.size = Pt(8.5)
r_nota_tbl.font.italic = True

# Guardar documento
doc.save(target_path)
print(f"Documento generado exitosamente en: {target_path}")
